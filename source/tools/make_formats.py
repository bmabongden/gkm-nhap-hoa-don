#!/usr/bin/env python3
"""Tạo bản .xls (BIFF8), Excel XML 2003 và HTML từ các hóa đơn mẫu .xlsx để kiểm thử bộ đọc của app.
Cần: pip install openpyxl xlwt. Chạy: python3 make_formats.py <thư mục samples>"""
import sys, os, html
import openpyxl, xlwt
from xml.sax.saxutils import escape

src = sys.argv[1]
out = os.path.join(src, "dinh-dang-khac")
os.makedirs(out, exist_ok=True)
names = {"Mau1": "Mau 1.xlsx", "Mau2": "Mau 2.xlsx", "Mau3": "Mau 3.xlsx", "Mau4": "Mau 4 - Toa hang Tan Hoa Loi.xlsx"}

def cells(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and str(c.value).strip() != "":
                yield c.row - 1, c.column - 1, c.value

for key, fn in names.items():
    path = os.path.join(src, fn)
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Sheet1")
    for r, c, v in cells(path):
        ws.write(r, c, v)
    wb.save(os.path.join(out, key + ".xls"))

# Excel XML 2003 và HTML cho Mẫu 4 và Mẫu 3
for key in ("Mau3", "Mau4"):
    rows = {}
    for r, c, v in cells(os.path.join(src, names[key])):
        rows.setdefault(r, {})[c] = v
    parts = ['<?xml version="1.0"?>\n<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet ss:Name="Sheet1"><Table>']
    h = ['<html><head><meta charset="utf-8"></head><body><table border="1">']
    last = -1
    for r in sorted(rows):
        for _ in range(r - last - 1):
            parts.append('<Row/>'); h.append('<tr><td></td></tr>')
        last = r
        parts.append('<Row>'); h.append('<tr>')
        prev = -1
        for c in sorted(rows[r]):
            v = rows[r][c]
            idx = ' ss:Index="%d"' % (c + 1) if c != prev + 1 else ''
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                parts.append('<Cell%s><Data ss:Type="Number">%s</Data></Cell>' % (idx, repr(v)))
            else:
                parts.append('<Cell%s><Data ss:Type="String">%s</Data></Cell>' % (idx, escape(str(v))))
            for _ in range(c - prev - 1):
                h.append('<td></td>')
            h.append('<td>%s</td>' % html.escape(str(v)))
            prev = c
        parts.append('</Row>'); h.append('</tr>')
    parts.append('</Table></Worksheet></Workbook>')
    h.append('</table></body></html>')
    open(os.path.join(out, key + "_xml2003.xls"), "w", encoding="utf-8").write("".join(parts))
    open(os.path.join(out, key + "_html.xls"), "w", encoding="utf-8").write("".join(h))
print("ok", sorted(os.listdir(out)))
